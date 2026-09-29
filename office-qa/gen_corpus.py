#!/usr/bin/env python3
"""SG Office 2 -- generate the independent round-trip corpus (ADR 0016, Phase 1).

Writes one .docx, one .xlsx and one .pptx with python-docx / openpyxl /
python-pptx -- tools that are NOT our engine -- so a round trip through our
engine is a real fidelity test, not an identity check on its own output. Each
file carries a handful of features the round-trip harness then checks survive.

    gen_corpus.py OUTDIR

Copyright (C) 2026 Stained Glass OS contributors
SPDX-License-Identifier: AGPL-3.0-or-later
"""
import sys
import os


def docx_file(path):
    import docx
    from docx.shared import Pt
    d = docx.Document()
    d.add_heading("SG Office round-trip: Documents", level=1)
    p = d.add_paragraph("The quick brown fox ")
    p.add_run("jumps").bold = True
    p.add_run(" over the lazy dog.")
    d.add_heading("A list", level=2)
    for item in ("alpha", "beta", "gamma"):
        d.add_paragraph(item, style="List Bullet")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Name"
    t.cell(0, 1).text = "Value"
    t.cell(1, 0).text = "answer"
    t.cell(1, 1).text = "42"
    d.save(path)


def xlsx_file(path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "Item"
    ws["B1"] = "Qty"
    ws["A2"] = "widgets"
    ws["B2"] = 3
    ws["A3"] = "gadgets"
    ws["B3"] = 4
    ws["B4"] = "=SUM(B2:B3)"          # formula -> cached value checked
    ws["B5"] = "=TEXTJOIN(\"-\",TRUE,B2,B3)"
    ws["C2"] = 1234.5
    ws["C2"].number_format = "#,##0.00"  # number format checked
    ws.merge_cells("A6:B6")           # merged range checked
    ws["A6"] = "merged footer"
    wb.save(path)


def pptx_file(path):
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "SG Office round-trip: Presentations"
    s.placeholders[1].text = "A subtitle line"
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Bullets"
    tf = s2.placeholders[1].text_frame
    tf.text = "first"
    tf.add_paragraph().text = "second"
    tf.add_paragraph().text = "third"
    prs.save(path)


def main():
    outdir = sys.argv[1] if len(sys.argv) > 1 else "corpus"
    os.makedirs(outdir, exist_ok=True)
    docx_file(os.path.join(outdir, "src.docx"))
    xlsx_file(os.path.join(outdir, "src.xlsx"))
    pptx_file(os.path.join(outdir, "src.pptx"))
    print("wrote src.docx, src.xlsx, src.pptx to", outdir)


if __name__ == "__main__":
    main()
